from __future__ import annotations

from pathlib import Path

from fastapi.testclient import TestClient
from sqlalchemy import text

from app.contracts import RightsStatus, SourceType
from app.db.models import Project
from app.main import create_app
from app.settings.config import Settings


LOOPBACK_ORIGIN = "http://127.0.0.1:8765"


def test_local_folder_preview_returns_candidates_without_creating_chapters(
    settings: Settings,
    db_session,
    tmp_path: Path,
) -> None:
    project = Project(
        id="018f0000-0000-7000-8000-000000008001",
        title="Import Preview",
        slug="import-preview",
        source_type=SourceType.SELF_AUTHORED.value,
        rights_status=RightsStatus.CLEARED.value,
    )
    db_session.add(project)
    db_session.commit()
    folder = tmp_path / "folder-import"
    folder.mkdir()
    (folder / "chapter-1.txt").write_text("Chuong 1\nPreview only.", encoding="utf-8")
    client = TestClient(create_app(settings=settings, acquire_lock=False), base_url=LOOPBACK_ORIGIN)
    csrf = client.get("/api/security/bootstrap").json()["csrfToken"]

    response = client.post(
        f"/api/projects/{project.id}/chapters/import/preview",
        data={"kind": "LOCAL_FOLDER", "localFolderPath": str(folder)},
        headers={"Origin": LOOPBACK_ORIGIN, "X-CSRF-Token": csrf},
    )

    assert response.status_code == 200
    candidate = response.json()["candidates"][0]
    assert candidate["ordinal"] == 1
    assert candidate["title"] == "Chuong 1"
    assert candidate["text"].splitlines() == ["Chuong 1", "Preview only."]
    assert candidate["sourcePath"] == "chapter-1.txt"
    assert candidate["warnings"] == []
    assert db_session.execute(text("SELECT COUNT(*) FROM chapters")).scalar_one() == 0


def test_docx_and_epub_preview_routes_return_candidates(settings: Settings, db_session, tmp_path: Path) -> None:
    from docx import Document

    project = Project(
        id="018f0000-0000-7000-8000-000000008101",
        title="Book Preview",
        slug="book-preview",
        source_type=SourceType.SELF_AUTHORED.value,
        rights_status=RightsStatus.CLEARED.value,
    )
    db_session.add(project)
    db_session.commit()
    docx_path = tmp_path / "book.docx"
    document = Document()
    document.add_paragraph("Chuong 1: Docx")
    document.add_paragraph("Docx body.")
    document.save(docx_path)
    epub_path = _write_minimal_epub(tmp_path / "book.epub")
    client = TestClient(create_app(settings=settings, acquire_lock=False), base_url=LOOPBACK_ORIGIN)
    csrf = client.get("/api/security/bootstrap").json()["csrfToken"]

    docx_response = client.post(
        f"/api/projects/{project.id}/chapters/import/preview",
        data={"kind": "DOCX"},
        files={"file": ("book.docx", docx_path.read_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers={"Origin": LOOPBACK_ORIGIN, "X-CSRF-Token": csrf},
    )
    epub_response = client.post(
        f"/api/projects/{project.id}/chapters/import/preview",
        data={"kind": "EPUB"},
        files={"file": ("book.epub", epub_path.read_bytes(), "application/epub+zip")},
        headers={"Origin": LOOPBACK_ORIGIN, "X-CSRF-Token": csrf},
    )

    assert docx_response.status_code == 200
    assert docx_response.json()["candidates"][0]["title"] == "Chuong 1: Docx"
    assert epub_response.status_code == 200
    assert epub_response.json()["candidates"][0]["title"] == "Chuong 1: Epub"
    assert db_session.execute(text("SELECT COUNT(*) FROM chapters")).scalar_one() == 0


def _write_minimal_epub(path: Path) -> Path:
    import zipfile

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("mimetype", b"application/epub+zip")
        archive.writestr(
            "META-INF/container.xml",
            b"""<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
<rootfiles><rootfile full-path="content.opf"/></rootfiles></container>""",
        )
        archive.writestr(
            "content.opf",
            b"""<package xmlns="http://www.idpf.org/2007/opf">
<manifest><item id="c1" href="chapter.xhtml" media-type="application/xhtml+xml"/></manifest>
<spine><itemref idref="c1"/></spine></package>""",
        )
        archive.writestr(
            "chapter.xhtml",
            b"""<html xmlns="http://www.w3.org/1999/xhtml"><body>
<h1>Chuong 1: Epub</h1><p>Epub body.</p>
</body></html>""",
        )
    return path
