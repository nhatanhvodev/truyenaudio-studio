from __future__ import annotations

import os
from pathlib import Path
import zipfile

import pytest

from app.modules.sources.archive_guard import InputArchiveTooLarge, InputArchiveUnsafe, inspect_archive
from app.modules.sources.docx import read_docx
from app.modules.sources.epub import read_epub
from app.modules.sources.folder import InputPathUnsafe, read_folder


@pytest.fixture
def zip_factory(tmp_path: Path):
    def build(members: dict[str, bytes]) -> Path:
        archive_path = tmp_path / "source.zip"
        with zipfile.ZipFile(archive_path, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, payload in members.items():
                archive.writestr(name, payload)
        return archive_path

    return build


@pytest.mark.parametrize("name", ["../escape.txt", "/absolute.txt", "C:/evil.txt"])
def test_archive_rejects_unsafe_member(tmp_path: Path, zip_factory, name: str) -> None:
    with pytest.raises(InputArchiveUnsafe):
        inspect_archive(zip_factory({name: b"x"}))

    assert not (tmp_path.parent / "escape.txt").exists()
    assert sorted(path.name for path in tmp_path.iterdir()) == ["source.zip"]


def test_archive_rejects_script_macro_executable_and_symlink_members(tmp_path: Path, zip_factory) -> None:
    with pytest.raises(InputArchiveUnsafe):
        inspect_archive(zip_factory({"word/vbaProject.bin": b"macro"}))
    with pytest.raises(InputArchiveUnsafe):
        inspect_archive(zip_factory({"EPUB/chapter.js": b"alert(1)"}))

    symlink_archive = tmp_path / "symlink.zip"
    symlink_info = zipfile.ZipInfo("link.txt")
    symlink_info.external_attr = 0o120777 << 16
    with zipfile.ZipFile(symlink_archive, "w") as archive:
        archive.writestr(symlink_info, "../outside.txt")

    with pytest.raises(InputArchiveUnsafe):
        inspect_archive(symlink_archive)


def test_zip_bomb_ratio_and_uncompressed_limit_are_enforced(monkeypatch: pytest.MonkeyPatch, zip_factory) -> None:
    with pytest.raises(InputArchiveTooLarge):
        inspect_archive(zip_factory({"chapter.txt": b"0" * 1024 * 1024}))

    monkeypatch.setattr("app.modules.sources.archive_guard.MAX_UNCOMPRESSED_BYTES", 4)
    with pytest.raises(InputArchiveTooLarge):
        inspect_archive(zip_factory({"chapter.txt": b"12345"}))


def test_read_epub_returns_spine_candidates_without_extracting(tmp_path: Path) -> None:
    epub_path = tmp_path / "book.epub"
    with zipfile.ZipFile(epub_path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("mimetype", b"application/epub+zip")
        archive.writestr(
            "META-INF/container.xml",
            b"""<?xml version="1.0"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles>
</container>""",
        )
        archive.writestr(
            "OEBPS/content.opf",
            b"""<?xml version="1.0"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0">
  <manifest>
    <item id="c1" href="chapters/ch1.xhtml" media-type="application/xhtml+xml"/>
    <item id="c2" href="chapters/ch2.xhtml" media-type="application/xhtml+xml"/>
  </manifest>
  <spine><itemref idref="c1"/><itemref idref="c2"/></spine>
</package>""",
        )
        archive.writestr(
            "OEBPS/chapters/ch1.xhtml",
            """<html xmlns="http://www.w3.org/1999/xhtml"><body>
<h1>第1章 Start</h1><p>Hello one.</p>
</body></html>""".encode("utf-8"),
        )
        archive.writestr(
            "OEBPS/chapters/ch2.xhtml",
            b"""<html xmlns="http://www.w3.org/1999/xhtml"><body>
<h1>Chuong 2: Tieptheo</h1><p>Hello two.</p>
</body></html>""",
        )

    candidates = read_epub(epub_path)

    assert [candidate.ordinal for candidate in candidates] == [1, 2]
    assert [candidate.title for candidate in candidates] == ["第1章 Start", "Chuong 2: Tieptheo"]
    assert "Hello one." in candidates[0].text
    assert sorted(path.name for path in tmp_path.iterdir()) == ["book.epub"]


def test_read_epub_rejects_inline_script(zip_factory) -> None:
    archive_path = zip_factory(
        {
            "mimetype": b"application/epub+zip",
            "META-INF/container.xml": b"""<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
<rootfiles><rootfile full-path="content.opf"/></rootfiles></container>""",
            "content.opf": b"""<package xmlns="http://www.idpf.org/2007/opf">
<manifest><item id="c1" href="ch1.xhtml" media-type="application/xhtml+xml"/></manifest>
<spine><itemref idref="c1"/></spine></package>""",
            "ch1.xhtml": b"""<html xmlns="http://www.w3.org/1999/xhtml"><body><script>bad()</script></body></html>""",
        }
    )

    with pytest.raises(InputArchiveUnsafe):
        read_epub(archive_path)


def test_read_docx_returns_paragraph_and_table_candidates(tmp_path: Path) -> None:
    from docx import Document

    docx_path = tmp_path / "book.docx"
    document = Document()
    document.add_paragraph("Chuong 1: Mo dau")
    document.add_paragraph("Doan mot.")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Thoai trong bang."
    document.add_paragraph("Chuong 2")
    document.add_paragraph("Doan hai.")
    document.save(docx_path)

    candidates = read_docx(docx_path)

    assert [candidate.ordinal for candidate in candidates] == [1, 2]
    assert candidates[0].title == "Chuong 1: Mo dau"
    assert "Thoai trong bang." in candidates[0].text
    assert "Doan hai." in candidates[1].text


def test_folder_reads_txt_in_natural_order_and_reports_preview_warnings(tmp_path: Path) -> None:
    root = tmp_path / "imports"
    root.mkdir()
    (root / "chapter-10.txt").write_text("Chuong 10\nMuoi.", encoding="utf-8")
    (root / "chapter-2.txt").write_text("Chuong 2\nHai.", encoding="utf-8")
    (root / "notes.md").write_text("ignore me", encoding="utf-8")
    (root / "empty.txt").write_text("", encoding="utf-8")

    candidates = read_folder(root)

    assert [candidate.source_path for candidate in candidates] == ["chapter-2.txt", "chapter-10.txt", "empty.txt"]
    assert [candidate.ordinal for candidate in candidates[:2]] == [2, 10]
    assert "EMPTY_CHAPTER" in candidates[2].warnings
    assert "ORDINAL_MISSING" in candidates[2].warnings


def test_folder_rejects_subfolder_escape(tmp_path: Path) -> None:
    root = tmp_path / "imports"
    outside = tmp_path / "outside"
    root.mkdir()
    outside.mkdir()
    (outside / "chapter.txt").write_text("Chuong 1\nOutside.", encoding="utf-8")

    with pytest.raises(InputPathUnsafe):
        read_folder(root, Path("..") / "outside")

    assert (outside / "chapter.txt").read_text(encoding="utf-8") == "Chuong 1\nOutside."


def test_folder_does_not_follow_symlink_or_reparse_point(tmp_path: Path) -> None:
    root = tmp_path / "imports"
    outside = tmp_path / "outside"
    root.mkdir()
    outside.mkdir()
    (outside / "chapter-1.txt").write_text("Chuong 1\nOutside.", encoding="utf-8")

    link_path = root / "linked.txt"
    try:
        os.symlink(outside / "chapter-1.txt", link_path)
    except OSError as exc:
        pytest.skip(f"symlink creation unavailable: {exc}")

    with pytest.raises(InputPathUnsafe):
        read_folder(root)

    assert (outside / "chapter-1.txt").read_text(encoding="utf-8") == "Chuong 1\nOutside."
